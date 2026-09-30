"""Create downloadable files from text/markdown: pdf, docx, pptx, xlsx, html, md, txt.

One entry point - create(content, fmt, out_path). Each format lazily imports its
library; the office formats (docx/pptx/xlsx) need the optional extra:
    pip install -e ".[docs]"
PDF (reportlab) and html/md/txt have no extra dependency.
"""

from __future__ import annotations

import html as _html
import re
from pathlib import Path

SUPPORTED = ("pdf", "docx", "pptx", "xlsx", "html", "md", "txt")


def create(content: str, fmt: str, out_path: Path, *, title: str = "") -> Path:
    """Write `content` to `out_path` in `fmt`. Returns the path. Raises ValueError
    on an unsupported format, ImportError if an optional library is missing."""
    fmt = fmt.lower().lstrip(".")
    if fmt not in SUPPORTED:
        raise ValueError(f"Unsupported format '{fmt}'. Supported: {', '.join(SUPPORTED)}")
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    {
        "pdf": _pdf,
        "docx": _docx,
        "pptx": _pptx,
        "xlsx": _xlsx,
        "html": _html_doc,
        "md": _raw,
        "txt": _raw,
    }[fmt](content, out_path, title)
    return out_path


def _raw(content: str, out: Path, title: str) -> None:
    out.write_text(content)


def _pdf(content: str, out: Path, title: str) -> None:
    from .writer import markdown_to_pdf

    markdown_to_pdf(content, out, title=title)


def _docx(content: str, out: Path, title: str) -> None:
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError('Install office libs: pip install -e ".[docs]"') from e
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for line in content.splitlines():
        s = line.rstrip()
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(s.lstrip()[2:], style="List Bullet")
        elif s.strip():
            doc.add_paragraph(s)
    doc.save(str(out))


def _pptx(content: str, out: Path, title: str) -> None:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ImportError('Install office libs: pip install -e ".[docs]"') from e
    prs = Presentation()
    blank = prs.slide_layouts[1]  # title + content
    slides = _split_slides(content, title)
    for head, bullets in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.title.text = head
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = b
    prs.save(str(out))


def _split_slides(content: str, title: str):
    """Each '# heading' starts a new slide; following lines are bullets."""
    slides, head, bullets = [], title or "Slides", []
    for line in content.splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            if bullets or slides:
                slides.append((head, bullets))
            head, bullets = s[2:], []
        elif s.lstrip().startswith(("- ", "* ")):
            bullets.append(s.lstrip()[2:])
        elif s.strip():
            bullets.append(s.strip())
    slides.append((head, bullets))
    return slides


def _xlsx(content: str, out: Path, title: str) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise ImportError('Install office libs: pip install -e ".[docs]"') from e
    wb = Workbook()
    ws = wb.active
    if title:
        ws.title = title[:31]
    for row in _rows_from_content(content):
        ws.append(row)
    wb.save(str(out))


def _rows_from_content(content: str) -> list[list]:
    """Parse a markdown table (| a | b |), else CSV/TSV lines, into rows."""
    lines = [ln for ln in content.splitlines() if ln.strip()]
    rows = []
    if any("|" in ln for ln in lines):
        for ln in lines:
            if "|" not in ln:
                continue
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if cells and all(set(c) <= set("-: ") for c in cells):
                continue  # markdown separator row
            rows.append(cells)
    else:
        for ln in lines:
            sep = "\t" if "\t" in ln else ","
            rows.append([c.strip() for c in ln.split(sep)])
    return rows


def _html_doc(content: str, out: Path, title: str) -> None:
    out.write_text(_markdown_to_html(content, title))


def _markdown_to_html(md: str, title: str = "") -> str:
    """Minimal, dependency-free markdown → a clean standalone HTML document."""
    body, in_list, in_code = [], False, False
    for line in md.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                body.append("</code></pre>")
                in_code = False
            else:
                body.append("<pre><code>")
                in_code = True
            continue
        if in_code:
            body.append(_html.escape(line))
            continue
        s = line.rstrip()
        if s.lstrip().startswith(("- ", "* ")):
            if not in_list:
                body.append("<ul>")
                in_list = True
            body.append(f"<li>{_inline(s.lstrip()[2:])}</li>")
            continue
        if in_list:
            body.append("</ul>")
            in_list = False
        if s.startswith("### "):
            body.append(f"<h3>{_inline(s[4:])}</h3>")
        elif s.startswith("## "):
            body.append(f"<h2>{_inline(s[3:])}</h2>")
        elif s.startswith("# "):
            body.append(f"<h1>{_inline(s[2:])}</h1>")
        elif s.strip():
            body.append(f"<p>{_inline(s)}</p>")
    if in_list:
        body.append("</ul>")
    head_title = _html.escape(title or "Document")
    return (
        '<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
        f"<title>{head_title}</title><style>"
        "body{font-family:system-ui,-apple-system,sans-serif;max-width:760px;"
        "margin:40px auto;padding:0 20px;line-height:1.6;color:#1a110a}"
        "pre{background:#1a110a;color:#f0dfc0;padding:14px;border-radius:8px;overflow:auto}"
        "code{font-family:ui-monospace,monospace}h1,h2,h3{line-height:1.25}"
        "</style></head><body>\n" + "\n".join(body) + "\n</body></html>\n"
    )


def _inline(text: str) -> str:
    """Escape, then apply **bold**, *italic*, and `code`."""
    t = _html.escape(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"\*(.+?)\*", r"<em>\1</em>", t)
    t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
    return t


# ── content preview for Office files (escaped; for the chat preview pane) ──────


def preview_html(path: Path, *, max_rows: int = 50, max_blocks: int = 200) -> str:
    """A safe (fully escaped) HTML fragment previewing a docx/pptx/xlsx file.

    Not pixel-perfect - a faithful *content* view using the same libs that wrote
    the file. Returns a small notice for unsupported types / missing libs.
    """
    ext = Path(path).suffix.lower().lstrip(".")
    try:
        if ext == "xlsx":
            return _preview_xlsx(path, max_rows)
        if ext == "docx":
            return _preview_docx(path, max_blocks)
        if ext == "pptx":
            return _preview_pptx(path)
    except ImportError:
        return '<div class="text-muted">Install the [docs] extra to preview Office files.</div>'
    except Exception as e:
        return f'<div class="text-muted">Preview unavailable: {_html.escape(str(e))}</div>'
    return '<div class="text-muted">No inline preview for this type - use the download link.</div>'


def _preview_xlsx(path: Path, max_rows: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            rows.append("<tr><td>… (truncated)</td></tr>")
            break
        tag = "th" if i == 0 else "td"
        cells = "".join(f"<{tag}>{_html.escape('' if c is None else str(c))}</{tag}>" for c in row)
        rows.append(f"<tr>{cells}</tr>")
    wb.close()
    return '<div class="file-preview-inner"><table>' + "".join(rows) + "</table></div>"


def _preview_docx(path: Path, max_blocks: int) -> str:
    from docx import Document

    parts = []
    for p in Document(str(path)).paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if "heading 1" in style:
            parts.append(f"<h3>{_html.escape(t)}</h3>")
        elif "heading" in style:
            parts.append(f"<h4>{_html.escape(t)}</h4>")
        elif "list" in style:
            parts.append(f"<li>{_html.escape(t)}</li>")
        else:
            parts.append(f"<p>{_html.escape(t)}</p>")
        if len(parts) >= max_blocks:
            parts.append("<p>… (truncated)</p>")
            break
    return '<div class="file-preview-inner">' + "".join(parts) + "</div>"


def _preview_pptx(path: Path) -> str:
    from pptx import Presentation

    parts = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        parts.append(f"<h4>Slide {i}</h4>")
        items = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = ("".join(r.text for r in para.runs) or para.text).strip()
                    if txt:
                        items.append(f"<li>{_html.escape(txt)}</li>")
        if items:
            parts.append("<ul>" + "".join(items) + "</ul>")
    return '<div class="file-preview-inner">' + "".join(parts) + "</div>"
